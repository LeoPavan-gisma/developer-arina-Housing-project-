"""Generate a detailed Word project report with charts and architecture diagrams."""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = ROOT / "docs" / "House_Price_Prediction_Project_Report.docx"
FIGURE_DIR = ROOT / "docs" / "report_figures"
INK = "17231F"
GREEN = "23654D"
ACID = "D9F16D"
MINT = "8BBE9B"
SLATE = "63736B"
PALE = "EDF2E8"
GRID = "D8DED5"
ACCENTS = [f"#{GREEN}", f"#{MINT}", f"#{ACID}", "#77918B"]


def _set_cell_shading(cell: Any, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _set_cell_margins(cell: Any, top: int = 100, start: int = 110,
                      bottom: int = 100, end: int = 110) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    properties.append(margins)


def _style_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, color in (("Title", 32, INK), ("Heading 1", 21, INK),
                              ("Heading 2", 14, GREEN), ("Heading 3", 10.5, INK)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(13 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("HOME VALUE STUDIO   /   PROJECT REPORT")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(SLATE)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("House Price Prediction  ·  September 2026  ·  ")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor.from_string(SLATE)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def _make_figures(data: pd.DataFrame, report: dict[str, Any]) -> dict[str, Path]:
    FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update({
        "font.family": "DejaVu Sans", "font.size": 9, "axes.titlesize": 12,
        "axes.labelcolor": f"#{INK}", "text.color": f"#{INK}",
        "axes.edgecolor": f"#{GRID}", "xtick.color": f"#{SLATE}",
        "ytick.color": f"#{SLATE}", "figure.facecolor": "white",
        "axes.facecolor": "white", "savefig.facecolor": "white",
    })
    figures: dict[str, Path] = {}

    fig, ax = plt.subplots(figsize=(8.2, 3.3))
    ax.hist(data["Price"] / 1_000_000, bins=18, color=f"#{GREEN}", edgecolor="white", linewidth=1.1)
    median = data["Price"].median() / 1_000_000
    ax.axvline(median, color="#B2762D", linewidth=2, linestyle="--", label=f"Median ₹{median:.1f}M")
    ax.set(title="Asking-price distribution", xlabel="Price (INR, millions)", ylabel="Number of listings")
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="y", color=f"#{GRID}", linewidth=.7)
    ax.legend(frameon=False)
    fig.tight_layout()
    figures["distribution"] = FIGURE_DIR / "price_distribution.png"
    fig.savefig(figures["distribution"], dpi=200, bbox_inches="tight")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(8.2, 4.0))
    location_colors = {name: ACCENTS[index % len(ACCENTS)] for index, name in enumerate(sorted(data["Location"].unique()))}
    for location, subset in data.groupby("Location"):
        ax.scatter(subset["Area"], subset["Price"] / 1_000_000, label=location,
                   color=location_colors[location], alpha=.78, s=38,
                   edgecolor="white", linewidth=.45)
    ax.set(title="Property size and asking price", xlabel="Area (square feet)", ylabel="Price (INR, millions)")
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(color=f"#{GRID}", linewidth=.65)
    ax.legend(title="Location", frameon=False, ncol=2)
    fig.tight_layout()
    figures["area_scatter"] = FIGURE_DIR / "area_vs_price.png"
    fig.savefig(figures["area_scatter"], dpi=200, bbox_inches="tight")
    plt.close(fig)

    summary = data.groupby("Location")["Price"].agg(["median", "count"]).sort_values("median")
    fig, ax = plt.subplots(figsize=(8.2, 3.6))
    bars = ax.barh(summary.index, summary["median"] / 1_000_000, color=ACCENTS[:len(summary)], height=.62)
    for bar, (_, row) in zip(bars, summary.iterrows(), strict=True):
        ax.text(bar.get_width() + .25, bar.get_y() + bar.get_height() / 2,
                f"₹{row['median'] / 1_000_000:.1f}M  ·  n={int(row['count'])}",
                va="center", fontsize=8, color=f"#{INK}")
    ax.set(title="Median asking price by location", xlabel="Median price (INR, millions)")
    ax.set_xlim(0, summary["median"].max() / 1_000_000 * 1.4)
    ax.spines[["top", "right", "left"]].set_visible(False)
    ax.tick_params(axis="y", length=0)
    ax.grid(axis="x", color=f"#{GRID}", linewidth=.65)
    fig.tight_layout()
    figures["location_median"] = FIGURE_DIR / "median_price_location.png"
    fig.savefig(figures["location_median"], dpi=200, bbox_inches="tight")
    plt.close(fig)

    names = list(report["algorithms"])
    fig, axes = plt.subplots(1, 3, figsize=(9.0, 3.6))
    panels = (("cv_r2_mean", "Cross-validation R²", "score"),
              ("r2", "Holdout R²", "score"), ("mae", "Holdout MAE", "currency"))
    for axis, (key, title, kind) in zip(axes, panels, strict=True):
        values = [report["algorithms"][name][key] for name in names]
        display = np.asarray(values) / 1_000_000 if kind == "currency" else np.asarray(values)
        bars = axis.bar(range(len(names)), display, color=[f"#{GREEN}", f"#{MINT}", f"#{ACID}"])
        axis.set_title(title, loc="left", fontweight="bold", fontsize=10)
        axis.set_xticks(range(len(names)), ["Linear\nRegression", "Random\nForest", "Gradient\nBoosting"], fontsize=7)
        axis.spines[["top", "right"]].set_visible(False)
        axis.grid(axis="y", color=f"#{GRID}", linewidth=.6)
        axis.set_axisbelow(True)
        for bar, value in zip(bars, values, strict=True):
            label = f"{value / 1_000_000:.2f}M" if kind == "currency" else f"{value:.3f}"
            axis.text(bar.get_x() + bar.get_width() / 2, bar.get_height(), label,
                      ha="center", va="bottom", fontsize=7, fontweight="bold")
        if kind == "score":
            axis.set_ylim(0, 1.12)
    fig.suptitle("Algorithm comparison: higher R² is better; lower MAE is better", x=.02, ha="left", fontweight="bold", fontsize=12)
    fig.tight_layout(rect=(0, 0, 1, .90))
    figures["model_comparison"] = FIGURE_DIR / "model_comparison.png"
    fig.savefig(figures["model_comparison"], dpi=200, bbox_inches="tight")
    plt.close(fig)

    importance = pd.DataFrame(report["feature_importance"]).sort_values("importance_mean")
    fig, ax = plt.subplots(figsize=(8.2, 3.7))
    bars = ax.barh(importance["feature"].str.replace("_", " "), importance["importance_mean"],
                   xerr=importance["importance_std"], color=ACCENTS[0], alpha=.92, capsize=3)
    ax.axvline(0, color=f"#{INK}", linewidth=.8)
    ax.set(title="Held-out permutation importance", xlabel="Decrease in R² after shuffling")
    ax.spines[["top", "right", "left"]].set_visible(False)
    ax.tick_params(axis="y", length=0)
    ax.grid(axis="x", color=f"#{GRID}", linewidth=.65)
    for bar in bars:
        ax.text(max(0, bar.get_width()) + .02, bar.get_y() + bar.get_height() / 2,
                f"{bar.get_width():.3f}", va="center", fontsize=8)
    fig.tight_layout()
    figures["importance"] = FIGURE_DIR / "feature_importance.png"
    fig.savefig(figures["importance"], dpi=200, bbox_inches="tight")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(9, 3.3))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 4)
    ax.axis("off")
    boxes = [
        (0.2, 2.3, 1.75, 1.1, "SOURCE DATA", "300 property listings", "#EDF2E8"),
        (2.25, 2.3, 1.75, 1.1, "PREPROCESS", "Validate · encode", "#E3EEE6"),
        (4.3, 2.3, 1.75, 1.1, "MODEL PIPELINE", "3 regressors · CV", "#DCEBDD"),
        (6.35, 2.3, 1.55, 1.1, "ARTIFACTS", "Versioned model", "#E8F0D3"),
        (8.2, 2.3, 1.6, 1.1, "SERVING", "Web UI · API", "#D9F16D"),
    ]
    for x, y, width, height, title, description, color in boxes:
        patch = FancyBboxPatch((x, y), width, height, boxstyle="round,pad=.08,rounding_size=.12",
                               linewidth=1, edgecolor=f"#{GREEN}", facecolor=color)
        ax.add_patch(patch)
        ax.text(x + width / 2, y + .73, title, ha="center", va="center", fontsize=8, fontweight="bold", color=f"#{INK}")
        ax.text(x + width / 2, y + .36, description, ha="center", va="center", fontsize=7, color=f"#{SLATE}")
    for start, end in ((1.97, 2.22), (4.02, 4.27), (6.07, 6.32), (7.92, 8.17)):
        ax.add_patch(FancyArrowPatch((start, 2.85), (end, 2.85), arrowstyle="-|>", mutation_scale=12,
                                     linewidth=1.3, color=f"#{GREEN}"))
    ax.text(5, 1.55, "Train/evaluate: held-out test set  +  5-fold cross-validation",
            ha="center", fontsize=9, color=f"#{GREEN}", fontweight="bold")
    ax.text(5, .82, "Inference: validate input → load latest manifest → predict → return estimate and empirical range",
            ha="center", fontsize=8, color=f"#{SLATE}")
    figures["architecture"] = FIGURE_DIR / "system_architecture.png"
    fig.savefig(figures["architecture"], dpi=220, bbox_inches="tight")
    plt.close(fig)
    return figures


def _paragraph(document: Document, text: str, bold_prefix: str | None = None) -> None:
    paragraph = document.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        paragraph.add_run(bold_prefix).bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)


def _bullet(document: Document, text: str) -> None:
    document.add_paragraph(text, style="List Bullet")


def _caption(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(10)
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(SLATE)


def _add_figure(document: Document, path: Path, caption: str, width: float = 6.55) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    _caption(document, caption)


def _add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        _set_cell_shading(cell, INK)
        _set_cell_margins(cell)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(8)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cells[index])
            if row_index % 2 == 1:
                _set_cell_shading(cells[index], "F1F4EF")
            for paragraph in cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(1)
                for run in paragraph.runs:
                    run.font.size = Pt(8)
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths, strict=True):
                cell.width = Inches(width)


def _add_page_break(document: Document) -> None:
    document.add_page_break()


def _format_inr(value: float) -> str:
    return f"₹{value:,.0f}"


def _make_document(data: pd.DataFrame, report: dict[str, Any], figures: dict[str, Path]) -> Document:
    document = Document()
    _style_document(document)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(38)
    run = paragraph.add_run("END-TO-END MACHINE LEARNING  /  2026")
    run.font.name = "Aptos"
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(GREEN)

    title = document.add_paragraph(style="Title")
    title.paragraph_format.space_before = Pt(15)
    title.add_run("House Price\nPrediction System")
    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(5)
    run = subtitle.add_run("A complete path from property data to a deployed valuation experience")
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor.from_string(SLATE)

    _add_figure(document, figures["architecture"], "Figure 1. Project architecture from source listings through training to web and API inference.", 6.7)
    summary_table = document.add_table(rows=1, cols=3)
    summary_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cover_metrics = (("300", "SOURCE LISTINGS"), ("3", "ALGORITHMS TESTED"), (f"{report['selected_metrics']['r2']:.3f}", "SELECTED MODEL TEST R²"))
    for cell, (value, label) in zip(summary_table.rows[0].cells, cover_metrics, strict=True):
        _set_cell_shading(cell, INK)
        _set_cell_margins(cell, top=180, bottom=180, start=150, end=150)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        value_run = p.add_run(value + "\n")
        value_run.font.size = Pt(19)
        value_run.font.bold = True
        value_run.font.color.rgb = RGBColor.from_string(ACID)
        label_run = p.add_run(label)
        label_run.font.size = Pt(7)
        label_run.font.bold = True
        label_run.font.color.rgb = RGBColor(235, 240, 230)

    document.add_paragraph("Prepared from the supplied house_prices.csv dataset", style="Subtitle")
    _paragraph(document, f"Model artifact: {report['model_version']}    |    Report generated: {datetime.now().strftime('%d %B %Y')}")
    _paragraph(document, "Currency values are shown in Indian rupees (INR). This is a learning project and not a professional valuation.")
    _add_page_break(document)

    document.add_heading("Contents", level=1)
    for entry in (
        "1. Executive summary", "2. Business problem and system goals", "3. Dataset and exploration",
        "4. Data preparation and feature engineering", "5. Algorithms and evaluation design",
        "6. Results and interpretation", "7. Deployment, API, and model versioning",
        "8. Testing, limitations, and responsible use", "9. Recommendations and next steps",
    ):
        _bullet(document, entry)
    document.add_heading("1. Executive Summary", level=1)
    _paragraph(document, "This project implements a modular machine-learning solution that estimates asking prices from property size, room counts, age, location, and property type. The workflow covers data validation, leakage-aware preprocessing, three-model comparison, cross-validation, holdout evaluation, feature interpretation, model persistence, REST serving, and an interactive Streamlit dashboard.")
    _paragraph(document, f"The supplied data contains {len(data)} rows and {len(data.columns)} source columns. An 80/20 split produced {report['training_rows']} training records and {report['test_rows']} final test records. The selected estimator is {report['selected_algorithm']}, selected by mean five-fold cross-validation R² ({report['selected_metrics']['cv_r2_mean']:.3f} ± {report['selected_metrics']['cv_r2_std']:.3f}). On the held-out test set it achieved R² {report['selected_metrics']['r2']:.3f}, MAE {_format_inr(report['selected_metrics']['mae'])}, RMSE {_format_inr(report['selected_metrics']['rmse'])}, and MAPE {report['selected_metrics']['mape']:.1f}%.")
    _paragraph(document, "A useful nuance: Random Forest scored the strongest single holdout R² in this run, while Gradient Boosting had the strongest mean CV R². The project follows its pre-declared selection rule and chooses Gradient Boosting; the holdout is reported as an independent check, not reused for selection.")

    document.add_heading("2. Business Problem and System Goals", level=1)
    document.add_heading("Problem statement", level=2)
    _paragraph(document, "A homeowner, agent, or analyst needs a quick initial price reference before a more detailed appraisal. Manual comparisons are time-consuming and inconsistent. The system provides a repeatable estimate based on the patterns available in a small listing dataset, together with a rough historical-error band and comparisons to the dataset's location/property-type median.")
    document.add_heading("Users and intended value", level=2)
    _bullet(document, "Property researchers can test how combinations of property attributes map to the fitted model's estimate.")
    _bullet(document, "Business stakeholders can compare algorithm performance and inspect which input groups the model relies on.")
    _bullet(document, "Developers can use the REST API and versioned pipeline as a compact model-serving example.")
    document.add_heading("Success criteria", level=2)
    _add_table(document, ["Area", "Success criterion", "Implemented evidence"], [
        ["Data", "Load and validate the supplied dataset", "300 rows checked; required columns and positive target enforced"],
        ["Modeling", "Compare at least three algorithms", "Linear Regression, Random Forest, Gradient Boosting"],
        ["Evaluation", "Use multiple metrics and validation", "MAE, RMSE, R², MAPE, five-fold CV"],
        ["Interpretation", "Explain influential inputs", "Held-out permutation importance"],
        ["Serving", "Support browser and API predictions", "Streamlit interface and FastAPI routes"],
        ["Operations", "Persist and identify model versions", "Timestamped joblib/report and latest manifest"],
    ])

    _add_page_break(document)
    document.add_heading("3. Dataset and Exploration", level=1)
    _paragraph(document, "The source file contains 300 property records. The target is Price, in INR. Property_ID is a record identifier, not a property characteristic, and is deliberately excluded from the model to prevent arbitrary identifier patterns from influencing predictions.")
    numeric_summary = data[["Area", "Bedrooms", "Bathrooms", "Age", "Price"]].describe().round(1)
    _add_table(document, ["Field", "Mean", "Std. dev.", "Minimum", "Median", "Maximum"], [
        [name, f"{numeric_summary.loc['mean', name]:,.1f}", f"{numeric_summary.loc['std', name]:,.1f}",
         f"{numeric_summary.loc['min', name]:,.0f}", f"{numeric_summary.loc['50%', name]:,.0f}",
         f"{numeric_summary.loc['max', name]:,.0f}"]
        for name in ("Area", "Bedrooms", "Bathrooms", "Age", "Price")
    ])
    _paragraph(document, f"The overall median asking price is {_format_inr(data['Price'].median())}; median area is {data['Area'].median():,.0f} square feet. Every required field is complete in the supplied CSV. These figures describe this sample only, not the current market across India.")
    _add_figure(document, figures["distribution"], "Figure 2. Distribution of asking prices. The dashed line marks the sample median.")
    _add_figure(document, figures["area_scatter"], "Figure 3. Area-price relationship, with points colored by location.")
    _add_figure(document, figures["location_median"], "Figure 4. Median asking price by location; n gives the sample count in each group.")

    _add_page_break(document)
    document.add_heading("4. Data Preparation and Feature Engineering", level=1)
    document.add_heading("Input fields", level=2)
    _add_table(document, ["Field", "Role", "Treatment"], [
        ["Property_ID", "Identifier", "Excluded from predictors"],
        ["Area", "Numeric predictor", "Median imputation, standard scaling"],
        ["Bedrooms / Bathrooms / Age", "Numeric predictors", "Median imputation, standard scaling"],
        ["Location", "Categorical predictor", "Most-frequent imputation, one-hot encoding"],
        ["Property_Type", "Categorical predictor", "Most-frequent imputation, one-hot encoding"],
        ["Price", "Target", "Positive numeric value in INR"],
    ])
    document.add_heading("Leakage-aware pipeline", level=2)
    _paragraph(document, "The data is split before fitting learned transformations. Numeric median imputers and scalers, as well as categorical imputers and one-hot category vocabularies, live inside a scikit-learn Pipeline/ColumnTransformer. During cross-validation each fold fits its preprocessing only on that fold's training portion. The final model is fitted on the 240-row training partition and evaluated once on the 60-row held-out partition.")
    document.add_heading("Prediction-time validation", level=2)
    _paragraph(document, "The inference layer requires all six predictors, converts numeric inputs consistently, rejects booleans and non-finite/out-of-range values, checks integer room/age fields, and requires non-empty categorical strings. Supported numeric bounds are 100–10,000 sq ft, 1–10 bedrooms, 1–10 bathrooms, and 0–150 years. One-hot encoding ignores categories not seen in training rather than crashing, while the interactive UI offers categories present in the dataset.")
    document.add_heading("Engineered representation", level=2)
    _paragraph(document, "No handcrafted ratios or geographic lookups are added because the dataset has no coordinates, neighborhood identifiers, or reliable external metadata. The feature-engineering work is the reproducible pipeline transformation: scaled numeric inputs and one-hot categorical indicators. This avoids inventing information and keeps the interface aligned with the source schema.")

    document.add_heading("5. Algorithms and Evaluation Design", level=1)
    _add_table(document, ["Algorithm", "Strengths", "Trade-offs"], [
        ["Linear Regression", "Fast, compact baseline; coefficients can be inspected", "Captures only linear effects unless interactions are added"],
        ["Random Forest", "Ensemble of decorrelated trees; handles nonlinear splits", "Less direct global explanation; can extrapolate poorly"],
        ["Gradient Boosting", "Sequential trees can model residual structure efficiently", "Requires tuning; sensitive to data size and training choices"],
    ])
    _paragraph(document, "The split is reproducible with random seed 42. Five-fold shuffled K-fold CV is applied to training data. Model selection maximizes mean CV R². The test metrics are independent reporting metrics and do not drive this choice. MAE gives typical absolute currency error; RMSE penalizes larger misses more; R² describes variance explained relative to a mean baseline; MAPE expresses absolute error as a percentage, but may become unstable when true values approach zero.")
    _add_figure(document, figures["model_comparison"], "Figure 5. Cross-validation R², held-out R², and held-out MAE for all three regressors.")

    _add_page_break(document)
    document.add_heading("6. Results and Interpretation", level=1)
    algorithm_rows = []
    for name, metrics in report["algorithms"].items():
        selected_mark = "  SELECTED" if name == report["selected_algorithm"] else ""
        algorithm_rows.append([
            name + selected_mark,
            f"{metrics['cv_r2_mean']:.3f} ± {metrics['cv_r2_std']:.3f}",
            f"{metrics['r2']:.3f}",
            _format_inr(metrics["mae"]),
            _format_inr(metrics["rmse"]),
            f"{metrics['mape']:.1f}%",
        ])
    _add_table(document, ["Algorithm", "CV R² mean ± SD", "Test R²", "Test MAE", "Test RMSE", "Test MAPE"], algorithm_rows)
    _paragraph(document, "Random Forest produced the lowest test MAE and RMSE and highest holdout R² in this particular split. Gradient Boosting achieved the highest mean CV R² and is therefore the persisted model under the configured selection criterion. This difference is a useful model-selection lesson: a single split can rank candidates differently from repeated validation. Small data makes all such estimates uncertain.")
    _add_figure(document, figures["importance"], "Figure 6. Permutation importance on the 60-row held-out set. Error bars show variation across shuffles.")
    importance = report["feature_importance"]
    leader = importance[0]
    second = importance[1]
    _paragraph(document, f"Area ranks first: shuffling it reduced test-set R² by {leader['importance_mean']:.2f} on average in this calculation. Location ranks second at {second['importance_mean']:.2f}. Bedroom count has a smaller measured contribution, while bathroom count and property type add little conditional predictive value for this fitted model and sample.")
    _paragraph(document, "Importance is not a percentage share of price, an estimated rupee effect, or a causal relationship. Correlated inputs can share or mask importance. A negative near-zero score is consistent with a feature contributing little beyond noise under the current test sample.")
    document.add_heading("Prediction-range interpretation", level=2)
    _paragraph(document, f"The interface reports estimate ± {_format_inr(report['absolute_error_90th_percentile'])}, the 90th percentile of absolute residuals in the holdout set. This is a transparent empirical error band, not a statistically calibrated confidence interval. With 60 test records, quantile uncertainty is substantial; future coverage is not guaranteed.")
    document.add_heading("Business implications", level=2)
    _bullet(document, "Use the estimate as a starting point for review, not as a final list price, appraisal, loan decision, or investment recommendation.")
    _bullet(document, "Area and location are the strongest information sources available in this table; data quality and location representativeness directly affect usefulness.")
    _bullet(document, "A group median is a sample comparison. The app does not control for precise neighborhood, condition, amenities, transaction date, or market liquidity.")

    document.add_heading("7. Deployment, API, and Model Versioning", level=1)
    _paragraph(document, "The Streamlit interface provides a property valuation form, sample market dashboard, model comparison, and session-level estimate history export. The FastAPI service exposes GET /health and POST /predict. Both serving paths call the same validation and inference functions, reducing behavioral differences between UI and API use.")
    _add_figure(document, figures["architecture"], "Figure 7. Runtime architecture and the shared training/inference path.", 6.7)
    document.add_heading("Prediction contract", level=2)
    _add_table(document, ["Request field", "Example", "Meaning"], [
        ["Area", "1500", "Area in square feet"], ["Bedrooms", "3", "Bedroom count"],
        ["Bathrooms", "2", "Bathroom count"], ["Age", "5", "Property age in years"],
        ["Location", "City Center", "Location category"], ["Property_Type", "Apartment", "Property type category"],
    ])
    _paragraph(document, "A successful response returns estimated_price, estimated_range_low, estimated_range_high, range_method, model_version, and algorithm. Invalid inputs return HTTP 422; unavailable model artifacts return HTTP 503. OpenAPI documentation is served at /docs.")
    document.add_heading("Artifact management", level=2)
    _paragraph(document, f"Training writes a timestamped joblib pipeline ({report['model_file']}), a matching JSON evaluation report, and models/latest.json. The manifest points to the active model/report pair. Keep all three artifacts together during deployment. Joblib uses Python serialization and must never be used to load untrusted files.")
    _paragraph(document, "Run from the project root: install requirements.txt, train with python -m scripts.train, launch the UI with streamlit run app/streamlit_app.py, or launch the API with python -m uvicorn app.api:app --reload.")

    _add_page_break(document)
    document.add_heading("8. Testing, Limitations, and Responsible Use", level=1)
    _paragraph(document, "The project test suite contains nine tests covering loading the real dataset, exclusion of identifiers, valid input normalization, missing and malformed/out-of-range fields, three-model training, report/manifest persistence, and prediction response/range construction. The live API was also smoke-tested for successful health and prediction calls and a 422 invalid-input response. The browser prediction flow and responsive layout were exercised.")
    _add_table(document, ["Limitation", "Why it matters", "Practical next step"], [
        ["Only 300 records", "Split/CV scores have sampling uncertainty", "Collect substantially more representative records"],
        ["Listing prices, not verified transactions", "Asking values can differ from sale values", "Define a verified target and data provenance"],
        ["Sparse location categories", "Broad categories hide neighborhood differences", "Use privacy-safe geographic features and enough samples"],
        ["No condition/amenity/time features", "Important value drivers are unobserved", "Add documented, quality-controlled fields"],
        ["No drift or feedback monitoring", "Market relationships change over time", "Track residuals, segment metrics, and data freshness"],
        ["Empirical range is not calibrated", "Coverage can vary by group or time", "Use larger calibration data and evaluate coverage"],
    ])
    document.add_heading("Responsible deployment", level=2)
    _paragraph(document, "Before business use, establish lawful data collection and retention, explain limitations to users, audit errors across relevant property and location groups, and ensure estimates do not become the sole basis for consequential decisions. The local development server has no authentication or production hardening. Production use needs access controls, logging, monitoring, resource limits, model governance, and an incident/retraining process.")

    document.add_heading("9. Recommendations and Next Steps", level=1)
    for text in (
        "Expand and verify the dataset, record its source and timestamp, and distinguish asking prices from closed transactions.",
        "Use a time-aware validation strategy when historical transaction dates become available; keep a truly future period for evaluation.",
        "Test error and interval coverage by location, property type, and price band; investigate material disparities before deployment.",
        "Tune tree model parameters inside cross-validation and compare against a robust median/location baseline.",
        "Calibrate prediction intervals on a separate calibration sample and report empirical coverage with interval width.",
        "Add monitoring for feature distributions, prediction ranges, API failures, and newly observed labeled outcomes.",
    ):
        _bullet(document, text)
    document.add_heading("Project files", level=2)
    _add_table(document, ["Path", "Responsibility"], [
        ["src/data_preprocessing.py", "Schema checks, feature definitions, request validation"],
        ["src/model_training.py", "Pipelines, model comparison, evaluation, interpretation, persistence"],
        ["src/model_inference.py", "Load latest artifact and produce validated estimates"],
        ["app/streamlit_app.py", "Interactive valuation and analytics interface"],
        ["app/api.py", "FastAPI health and prediction endpoints"],
        ["scripts/train.py / create_report.py", "Training and report-generation commands"],
        ["tests/", "Focused regression tests for data, training, and inference"],
    ])
    closing = document.add_paragraph()
    closing.paragraph_format.space_before = Pt(16)
    closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = closing.add_run("END OF PROJECT REPORT")
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(GREEN)

    return document


def generate_report() -> Path:
    data = pd.read_csv(ROOT / "house_prices.csv")
    report = json.loads((ROOT / "models" / "latest.json").read_text(encoding="utf-8"))
    metrics_path = ROOT / "models" / report["report_file"]
    metrics = json.loads(metrics_path.read_text(encoding="utf-8"))
    figures = _make_figures(data, metrics)
    document = _make_document(data, metrics, figures)
    document.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    print(f"Created {generate_report()}")